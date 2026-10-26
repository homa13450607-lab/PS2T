from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.engines.schema import SpeechResult
from app.core.config import settings

class WordExporter:
    @staticmethod
    def export(result: SpeechResult, filename: str) -> Path:
        output_dir = settings.get_path("outputs_dir") / "docx"
        output_dir.mkdir(parents=True, exist_ok=True)
        file_path = output_dir / f"{filename}.docx"

        doc = Document()
        
        title = doc.add_heading(f"گزارش پیاده‌سازی صوت: {filename}", level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        p_meta = doc.add_paragraph()
        p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_meta.add_run(f"مدت زمان فایل: {result.audio_duration_seconds} ثانیه | ").bold = True
        p_meta.add_run(f"زمان پردازش: {result.processing_time_seconds} ثانیه | ")
        p_meta.add_run(f"سخت‌افزار: {result.device_used}")

        doc.add_heading("متن کامل", level=2).alignment = WD_ALIGN_PARAGRAPH.RIGHT
        # Keep paragraph boundaries from the formatted transcript in the DOCX.
        for paragraph_text in (result.full_text or "").split("\n\n"):
            if not paragraph_text.strip():
                continue
            p_full = doc.add_paragraph(paragraph_text.strip())
            p_full.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        doc.add_heading("جزئیات زمان‌بندی و میزان اطمینان", level=2).alignment = WD_ALIGN_PARAGRAPH.RIGHT
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'درصد اطمینان'
        hdr_cells[1].text = 'متن'
        hdr_cells[2].text = 'زمان پایان'
        hdr_cells[3].text = 'زمان شروع'

        for seg in result.segments:
            row_cells = table.add_row().cells
            row_cells[0].text = f"{seg.confidence}%"
            row_cells[1].text = seg.text
            row_cells[2].text = seg.end_time
            row_cells[3].text = seg.start_time

        doc.save(file_path)
        return file_path